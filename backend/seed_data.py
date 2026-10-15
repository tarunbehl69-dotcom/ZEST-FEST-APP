import os
from pathlib import Path
import openpyxl
from backend.database import get_db_connection, init_db

BASE_DIR = Path(__file__).resolve().parent.parent
DATA_PATH = BASE_DIR / "data" / "events_seed.xlsx"


def seed_events():
    """Seed the events table from data/events_seed.xlsx if the table is empty."""
    # Ensure database tables exist first
    init_db()

    conn = get_db_connection()
    cursor = conn.cursor()

    # Check if events table already has rows
    cursor.execute("SELECT COUNT(*) AS count FROM events")
    count = cursor.fetchone()["count"]
    if count > 0:
        print(f"Events table already has {count} rows. Skipping seed.")
        conn.close()
        return

    # Check if the Excel file exists
    if not DATA_PATH.exists():
        conn.close()
        raise FileNotFoundError(
            f"Seed file not found at {DATA_PATH}. "
            "Please copy the provided events_seed.xlsx file into the data/ folder before running this script."
        )

    # Load workbook and Events sheet
    wb = openpyxl.load_workbook(DATA_PATH)
    if "Events" not in wb.sheetnames:
        conn.close()
        raise ValueError(
            f"Sheet 'Events' not found in {DATA_PATH}. Available sheets: {wb.sheetnames}"
        )

    sheet = wb["Events"]

    # Verify header row
    expected_headers = [
        "id",
        "title",
        "category",
        "description",
        "event_date",
        "event_time",
        "venue",
        "capacity",
        "fee",
        "icon",
    ]

    header_row = [str(cell.value).strip() if cell.value is not None else "" for cell in sheet[1]]
    # Match first 10 headers
    if header_row[: len(expected_headers)] != expected_headers:
        conn.close()
        raise ValueError(
            f"Header row does not match expected format.\nExpected: {expected_headers}\nFound: {header_row[:len(expected_headers)]}"
        )

    inserted_count = 0
    # Read rows starting from row 2
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue

        (
            event_id,
            title,
            category,
            description,
            event_date,
            event_time,
            venue,
            capacity,
            fee,
            icon,
        ) = row[:10]

        cursor.execute(
            """
            INSERT INTO events (id, title, category, description, event_date, event_time, venue, capacity, fee, icon)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                int(event_id),
                str(title),
                str(category),
                str(description) if description is not None else "",
                str(event_date),
                str(event_time),
                str(venue),
                int(capacity),
                int(fee) if fee is not None else 0,
                str(icon) if icon is not None else "",
            ),
        )
        inserted_count += 1

    conn.commit()
    conn.close()
    print(f"Successfully seeded {inserted_count} events from {DATA_PATH}")


if __name__ == "__main__":
    seed_events()
    # Query and display table contents
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT id, title, category, event_date, event_time, venue, capacity, fee, icon FROM events ORDER BY id")
    rows = cursor.fetchall()
    print("\n--- Current Events in Database ---")
    for r in rows:
        print(f"[{r['id']}] {r['icon']} {r['title']} ({r['category']}) - {r['event_date']} {r['event_time']} @ {r['venue']} | Cap: {r['capacity']} | Fee: ₹{r['fee']}")
    conn.close()
